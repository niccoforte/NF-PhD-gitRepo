"""Digitise published Hsieh Fig. 10 from its native embedded raster.

This is graphical evidence, not author-provided numerical data. The published
figure is a JPEG, so marker centres are fitted by colour-specific inversion
symmetry, with threshold/window controls and explicit pixel bounds. Printed
density and manuscript-assigned density are both retained. No workbook, FEA,
shared calculation default or original source file is modified.
"""

import argparse
import csv
import hashlib
import importlib.metadata
import itertools
import json
import math
from pathlib import Path
import subprocess
import sys

import numpy as np
from PIL import Image
from scipy.ndimage import map_coordinates
from scipy.optimize import minimize


PUBLISHED_SHA256 = "fb75f47c1e0b3db57a2bc0fba87a2b03a93ba74efb817f4a66123af013ae56c0"
# Visual locator seeds only; the fitted coordinates and controls are exported.
# Figure symbols are bow-ties, six-pointed stars and hexagons, each ONE point.
POINT_SEEDS = {
    "tri": {"initial": [(660, 403), (724, 377), (789, 354)],
            "final": [(660, 326), (724, 303), (789, 280)]},
    "kagome": {"initial": [(660, 415), (724, 383), (789, 366)],
               "final": [(660, 350), (724, 329), (789, 305)]},
    "hex": {"initial": [(660, 638), (724, 596), (789, 559)],
            "final": [(660, 589), (724, 560), (789, 522)]},
}
RHO = (.13, .16, .20)
Y_ANCHORS = ((100., 201), (10., 456), (1., 711), (.1, 966),
             (.01, 1221), (.001, 1474))
X_ANCHORS = ((10., 126), (100., 822), (1000., 1518))


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def extract_figure(pdf, output, pdf_python):
    if digest(pdf) != PUBLISHED_SHA256:
        raise ValueError("Source differs from the inspected published edition")
    # The configured scientific environment need not carry PDF dependencies.
    # An explicit bundled PDF Python can extract the original image bytes.
    image_path = output / "fig10-source.jpg"
    extractor = """
import json, sys
from collections import Counter
from pathlib import Path
from pypdf import PdfReader
reader = PdfReader(sys.argv[1])
page = reader.pages[12]
if len(page.images) != 1:
    raise ValueError('Expected exactly one image on published page13')
image = page.images[0]
if image.image.format != 'JPEG':
    raise ValueError('Expected the original embedded JPEG')
Path(sys.argv[2]).write_bytes(image.data)
print(json.dumps({'page_count':len(reader.pages), 'page':13,
                  'image_count':1, 'image_name':image.name,
                  'native_size':list(image.image.size),
                  'page_operator_counts':dict(Counter(
                      op.decode('ascii') for _, op in page.get_contents().operations))}))
"""
    extraction = subprocess.run([str(pdf_python), "-B", "-c", extractor, str(pdf), str(image_path)],
                                check=True, capture_output=True, text=True)
    listing = json.loads(extraction.stdout)
    im = Image.open(image_path).convert("RGB")
    if im.size != (1854, 1625):
        raise ValueError("Unexpected native raster dimensions: " + repr(im.size))
    return image_path, np.asarray(im, dtype=float), listing


def line_centroid(counts, approximate, radius=12, border=False):
    lo, hi = approximate - radius, approximate + radius + 1
    local = counts[lo:hi]
    best = int(np.argmax(local)) + lo
    if border:
        selected = np.flatnonzero(local >= .95 * np.max(local)) + lo
        # The thick left plot border is the contiguous maximum near the seed.
        selected = selected[np.abs(selected - approximate) <= 5]
        return float(np.average(selected, weights=counts[selected]))
    positions = np.arange(best - 2, best + 3)
    weights = counts[positions].astype(float)
    weights[weights < .5 * np.max(weights)] = 0
    return float(np.average(positions, weights=weights))


def calibrate(array):
    gray = ((np.ptp(array, axis=2) < 16) & (array.mean(axis=2) < 230)
            & (array.mean(axis=2) > 35))
    horizontal = gray[:, 150:1820].sum(axis=1)
    vertical = gray[50:1450, :].sum(axis=0)
    dark = (np.ptp(array, axis=2) < 16) & (array.mean(axis=2) < 230)
    left_counts = dark[50:1450, :].sum(axis=0)
    ya = [{"value": value, "pixel": line_centroid(horizontal, seed), "seed": seed}
          for value, seed in Y_ANCHORS]
    xa = [{"value": value,
           "pixel": line_centroid(left_counts if i == 0 else vertical, seed, border=i == 0),
           "seed": seed} for i, (value, seed) in enumerate(X_ANCHORS)]
    result = {}
    for name, anchors in (("toughness", ya), ("printed_density", xa)):
        logs = np.log10([a["value"] for a in anchors])
        pixels = np.array([a["pixel"] for a in anchors])
        slope, intercept = np.polyfit(logs, pixels, 1)
        residuals = pixels - (slope * logs + intercept)
        loo = []
        for i in range(len(anchors)):
            keep = np.arange(len(anchors)) != i
            m, c = np.polyfit(logs[keep], pixels[keep], 1)
            loo.append(float((pixels[i] - c) / m - logs[i]))
        result[name] = {"anchors": anchors, "pixel_per_log10": float(slope),
                        "pixel_intercept": float(intercept),
                        "residual_pixels": residuals.tolist(),
                        "leave_one_out_log10_errors": loo,
                        "anchor_uncertainty_pixels": 1.0}
    return result


def colour_signal(array, topology, threshold):
    r, g, b = np.moveaxis(array, 2, 0)
    if topology == "tri":
        return np.clip((r - np.maximum(g, b) - threshold) / (255 - threshold), 0, 1)
    if topology == "kagome":
        return np.clip((threshold - np.maximum.reduce((r, g, b))) / threshold, 0, 1)
    return np.clip((b - np.maximum(r, g) - threshold) / (180 - threshold), 0, 1)


def symmetry_centre(signal, seed, radius):
    sx, sy = seed
    yy, xx = np.mgrid[sy-radius:sy+radius+1, sx-radius:sx+radius+1]
    observed = signal[yy, xx]
    weights = np.exp(-((xx-sx)**2 + (yy-sy)**2) / (2 * 17.**2))
    observed_norm = np.sum(weights * observed**2)

    def objective(centre):
        reflected = map_coordinates(signal, [2*centre[1]-yy, 2*centre[0]-xx],
                                    order=1, mode="constant", cval=0)
        denominator = np.sqrt(observed_norm * np.sum(weights * reflected**2))
        return -float(np.sum(weights * observed * reflected) / denominator) if denominator else 1.

    fit = minimize(objective, np.asarray(seed, float), method="Powell",
                   bounds=((sx-5, sx+5), (sy-5, sy+5)),
                   options={"xtol": .002, "ftol": 1e-8})
    if not fit.success or np.max(np.abs(fit.x - seed)) > 4.9:
        raise ValueError("Marker fit reached locator bound: " + repr((seed, fit)))
    return fit.x, -float(fit.fun)


def map_value(pixel, calibration):
    return float(10 ** ((pixel - calibration["pixel_intercept"])
                       / calibration["pixel_per_log10"]))


def graphical_bounds(pixel, marker_bound, calibration):
    """All corners of bounded anchor/marker perturbations; NOT a confidence interval."""
    anchors = calibration["anchors"]
    logs = np.log10([a["value"] for a in anchors])
    pixels = np.array([a["pixel"] for a in anchors])
    values = []
    for signs in itertools.product((-1., 1.), repeat=len(anchors)):
        m, c = np.polyfit(logs, pixels + np.array(signs), 1)
        for pp in (pixel-marker_bound, pixel+marker_bound):
            values.append(10 ** ((pp-c) / m))
    return float(min(values)), float(max(values))


def local_bracket_value(pixel, calibration):
    anchors = sorted(calibration["anchors"], key=lambda a: a["pixel"])
    for left, right in zip(anchors, anchors[1:]):
        if left["pixel"] <= pixel <= right["pixel"]:
            return float(10 ** np.interp(pixel, [left["pixel"], right["pixel"]],
                                        np.log10([left["value"], right["value"]])))
    raise ValueError("Digitised point lies outside calibrated tick range")


def digitise(array, calibration):
    rows, controls = [], []
    thresholds = {"tri": (45., 65., 85.), "kagome": (90., 120., 150.),
                  "hex": (30., 37.5, 45.)}
    for topology, stages in POINT_SEEDS.items():
        for stage, seeds in stages.items():
            for rho, seed in zip(RHO, seeds):
                fits = []
                for threshold in thresholds[topology]:
                    signal = colour_signal(array, topology, threshold)
                    for radius in (20, 23):
                        centre, score = symmetry_centre(signal, seed, radius)
                        fits.append({"threshold": threshold, "radius": radius,
                                     "x": float(centre[0]), "y": float(centre[1]),
                                     "symmetry_score": score})
                # Median over prespecified controls, never selected against p1 or Fig.8.
                x, y = np.median([[f["x"], f["y"]] for f in fits], axis=0)
                spread_x = max(abs(f["x"]-x) for f in fits)
                spread_y = max(abs(f["y"]-y) for f in fits)
                xb, yb = 1. + spread_x, 1. + spread_y
                k = map_value(y, calibration["toughness"])
                k_local = local_bracket_value(y, calibration["toughness"])
                klo, khi = graphical_bounds(y, yb, calibration["toughness"])
                density = map_value(x, calibration["printed_density"])
                dlo, dhi = graphical_bounds(x, xb, calibration["printed_density"])
                point_id = "{}-{}-{:02d}".format(topology, stage, round(100*rho))
                rows.append({"point_id": point_id, "topology": topology, "stage": stage,
                             "quantity": "K_JIC" if stage == "initial" else "K_JFC",
                             "nominal_relative_density": rho, "bar_length_mm": 3.,
                             "manuscript_density_kg_m3": 4500*rho,
                             "printed_axis_density_kg_m3": density,
                             "printed_density_graphical_lower": dlo,
                             "printed_density_graphical_upper": dhi,
                             "K_MPa_sqrt_m": k, "K_graphical_lower": klo,
                             "K_graphical_upper": khi,
                             "K_local_bracket_calibration_MPa_sqrt_m": k_local,
                             "centre_x_pixel": float(x),
                             "centre_y_pixel": float(y), "centre_bound_x_pixel": xb,
                             "centre_bound_y_pixel": yb, "source_page": 13,
                             "source_figure": "10", "source_type": "native_embedded_JPEG",
                             "minimum_symmetry_score": min(f["symmetry_score"] for f in fits),
                             "uncertainty_kind": "bounded_graphical_sensitivity_not_statistical_CI"})
                controls.append({"point_id": point_id, "locator_seed": seed, "fits": fits})
    return rows, controls


def plots(image_path, rows, output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    image = Image.open(image_path)
    colors = {"tri": "#d62728", "kagome": "#111111", "hex": "#3b55ab"}
    fig, ax = plt.subplots(figsize=(11.4, 10), constrained_layout=True)
    ax.imshow(image)
    for i, row in enumerate(rows, 1):
        x, y = row["centre_x_pixel"], row["centre_y_pixel"]
        ax.plot(x, y, "+", color="#00d5e8", ms=8, mew=1)
        ax.text(x+14, y, str(i), fontsize=6.7, color="#007887",
                bbox={"facecolor": "white", "alpha": .8, "pad": .3, "edgecolor": "none"})
    ax.set_title("Hsieh Fig. 10: fitted centres (cyan), IDs follow digitised CSV row order")
    ax.axis("off")
    fig.savefig(output / "fig10-overlay.png", dpi=200)
    plt.close(fig)

    fig, axes = plt.subplots(1, 3, figsize=(11, 8.5), constrained_layout=True)
    for ax, rho in zip(axes, RHO):
        xcentre = np.mean([r["centre_x_pixel"] for r in rows
                           if r["nominal_relative_density"] == rho])
        ax.imshow(image)
        ax.set_xlim(xcentre-34, xcentre+100)
        ax.set_ylim(665, 250)
        for i, row in enumerate(rows, 1):
            if row["nominal_relative_density"] != rho:
                continue
            x, y = row["centre_x_pixel"], row["centre_y_pixel"]
            ax.plot(x, y, "+", color="#00d5e8", ms=9, mew=1)
            ax.annotate("{}: {} {}\n{:.3f} MPa sqrt(m)".format(
                i, row["topology"], row["stage"], row["K_MPa_sqrt_m"]),
                (x, y), (xcentre+22, y + (-9 if row["topology"] == "tri" else 8)),
                fontsize=7, bbox={"facecolor": "white", "alpha": .95, "edgecolor": "none", "pad": 1},
                arrowprops={"arrowstyle": "-", "lw": .6, "color": "#007887"})
        ax.set_title("Nominal relative density {:.0%}".format(rho))
        ax.axis("off")
    fig.suptitle("Marker identity / centre verification; source pixels are unchanged")
    fig.savefig(output / "fig10-marker-controls.png", dpi=200)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(10, 4), constrained_layout=True)
    for ax, stage in zip(axes, ("initial", "final")):
        for topology in POINT_SEEDS:
            selected = [r for r in rows if r["topology"] == topology and r["stage"] == stage]
            values = np.array([r["K_MPa_sqrt_m"] for r in selected])
            errors = np.array([[r["K_MPa_sqrt_m"]-r["K_graphical_lower"] for r in selected],
                               [r["K_graphical_upper"]-r["K_MPa_sqrt_m"] for r in selected]])
            ax.errorbar([100*r["nominal_relative_density"] for r in selected], values,
                        yerr=errors, marker="o", color=colors[topology], capsize=3,
                        label=topology)
        ax.set(xlabel="Nominal relative density (%)", ylabel="K (MPa sqrt(m))",
               title="Initial toughness K_JIC" if stage == "initial" else "Final toughness K_JFC",
               xticks=[13, 16, 20])
        ax.grid(alpha=.25)
        ax.legend()
    fig.suptitle("Hsieh Fig. 10 graphical extraction, bar length 3 mm\n"
                 "Lines join markers; bars show bounded pixel sensitivity, not confidence intervals", fontsize=10)
    fig.savefig(output / "fig10-digitised.png", dpi=200)
    plt.close(fig)


def extract_fig8_workbook(workbook, output, pdf_python):
    """Preserve original literal digitisation, not transformed workbook columns."""
    extractor = """
import json,sys,hashlib
from pathlib import Path
from openpyxl import load_workbook
p=Path(sys.argv[1]);w=load_workbook(p,read_only=True,data_only=False);v=load_workbook(p,read_only=True,data_only=True)
s,c=w['Deshpande'],v['Deshpande'];rows=[]
for figure,indices in [('8a',range(15,33)),('8b',range(38,44))]:
 for i in indices:
  literal=s['F'+str(i)].value
  if not isinstance(literal,(int,float)):raise ValueError('Expected literal digitisation in F'+str(i))
  rows.append({'figure':figure,'workbook_row':i,'topology':c['A'+str(i)].value,
               'relative_density':c['B'+str(i)].value,'l_mm':1000*c['C'+str(i)].value,
               'literal_ordinate':literal,'source_cell':'Deshpande!F'+str(i),
               'length_formula':s['C'+str(i)].value})
material={'parent_UTS_workbook_MPa':c['C3'].value,'parent_yield_MPa':c['C4'].value,
          'parent_E_MPa':c['C2'].value,'failure_strain':c['B6'].value,
          'table1_RO_exponent':c['B7'].value}
w.close();v.close();print(json.dumps({'workbook':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'rows':rows,'material':material}))
"""
    result = subprocess.run([str(pdf_python), "-B", "-c", extractor, str(workbook)],
                            check=True, capture_output=True, text=True)
    captured = json.loads(result.stdout)
    rows = captured.pop("rows")
    for row in rows:
        row["ordinate_definition"] = "K_JIC/(sigma_TS*sqrt(l_m))" if row["figure"] == "8a" else "K_JIC/sigma_TS"
        row["ordinate_unit"] = "dimensionless" if row["figure"] == "8a" else "sqrt(m)"
        row["abscissa_definition"] = "relative_density" if row["figure"] == "8a" else "sqrt(l_mm)"
        row["abscissa_value"] = row["relative_density"] if row["figure"] == "8a" else math.sqrt(row["l_mm"])
        row["source_abscissa_printed_unit"] = "dimensionless" if row["figure"] == "8a" else "sqrt(m); tick positions correspond to sqrt(mm)"
        row["digitisation_uncertainty"] = "not recorded in original workbook"
        row["reference_identity"] = "{}-{}-rho{}".format(row["figure"], row["topology"], row["relative_density"]) + (
            "" if row["figure"] == "8a" else "-l{}".format(row["l_mm"]))
        row["duplicated_length_assignment"] = row["figure"] == "8a"
        row["sigma_TS_numeric_in_source"] = "not explicitly reported for titanium; workbook uses parent1111.79203277635MPa"
    with (output / "fig8-original-digitisation.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    captured.update(source="Original workbook F15:F32 and F38:F43, read-only",
                    published_pdf_sha256=PUBLISHED_SHA256, published_page=11,
                    published_axes={"8a":"K_JIC/(sigma_TS*sqrt(l)) vs relative density",
                                    "8b":"K_JIC/sigma_TS in sqrt(m) vs printed sqrt(l) in sqrt(m); numeric x ticks correspond to sqrt(mm)"},
                    fig8a_unique_points=9, fig8a_workbook_rows=18, fig8b_points=6,
                    limits=["The 3/10mm Fig8a rows repeat nine identical digitised ordinates; not independent points",
                            "Published axes say sigma_TS; Table1/Fig10 support a yield-stress interpretation only as a separate inference",
                            "Original Hsieh ordinates have not been renormalised or corrected",
                            "Fig8a/8b length-factor inconsistency remains unresolved"])
    (output / "fig8-source-provenance.json").write_text(json.dumps(captured, indent=2), encoding="utf-8")
    return rows


def run(args):
    pdf, output = Path(args.source_pdf).resolve(), Path(args.out).resolve()
    output.mkdir(parents=True, exist_ok=True)
    image_path, array, listing = extract_figure(pdf, output, Path(args.pdf_python))
    calibration = calibrate(array)
    rows, controls = digitise(array, calibration)
    with (output / "fig10-digitised.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    provenance = {"source_pdf": str(pdf), "source_sha256": digest(pdf),
                  "published_page": 13, "figure": 10, "image": image_path.name,
                  "image_sha256": digest(image_path), "image_size_pixels": [1854, 1625],
                  "pdf_image_inspection": listing, "pixel_origin": "top left; x right, y down",
                  "calibration": calibration, "marker_controls": controls,
                  "digitiser_script_sha256": digest(__file__),
                  "software_versions": {name: importlib.metadata.version(name)
                                        for name in ("numpy", "scipy", "Pillow", "matplotlib")},
                  "method": "Median colour-specific 180-degree symmetry fit over three thresholds and two windows",
                  "graphical_bounds": "All sign corners of +/-1 pixel calibration anchors, plus +/- (1 pixel + control spread) marker centres; not a statistical confidence interval",
                  "assignment": "Left/middle/right groups identified as nominal rho .13/.16/.20 from the manuscript; manuscript density uses its stated parent density4500 kg/m3",
                  "limits": ["Native source figure is raster; there are no recoverable vector marker centres",
                             "Printed density positions disagree by approximately tenfold with manuscript densities",
                             "Digitisation recovers plotted values, not unrounded simulation results",
                             "Final K_JFC points must not be used as initial K_JIC validation targets",
                             "Only marker coordinates are data; connecting dashed source lines add no numerical observations",
                             "No Fig.8 or p1 simulation values enter marker fits or calibration"]}
    (output / "fig10-provenance.json").write_text(json.dumps(provenance, indent=2), encoding="utf-8")
    plots(image_path, rows, output)
    if args.workbook:
        extract_fig8_workbook(Path(args.workbook).resolve(), output, Path(args.pdf_python))
    print(json.dumps({"points": len(rows), "csv": str(output / "fig10-digitised.csv"),
                      "max_marker_control_spread_pixels": max(
                          max(r["centre_bound_x_pixel"], r["centre_bound_y_pixel"])-1 for r in rows)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-pdf", required=True)
    parser.add_argument("--pdf-python", default=sys.executable,
                        help="Python with pypdf installed; can be the bundled PDF runtime")
    parser.add_argument("--out", default="Z:/p1/sims/Ti/Validation/audit/benchmark")
    parser.add_argument("--workbook", help="Read-only original Deshpande Fig8 literal digitisation")
    run(parser.parse_args())
